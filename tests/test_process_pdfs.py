import datetime

import openpyxl
import pytest

import process_pdfs as pp

TODAY = datetime.date.today()


def make_template(folder, name="xxxxx-SH-S-0001-Document Issue Sheet.xlsx", dates=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for offset, (d, m, y) in enumerate(dates):
        col = pp.FIRST_ISSUE_COL + offset
        ws.cell(row=3, column=col, value=d)
        ws.cell(row=4, column=col, value=m)
        ws.cell(row=5, column=col, value=y)
    wb.save(folder / name)
    return folder / name


def run(folder, number="12345", name="Test Project"):
    return pp.run(["--folder", str(folder), "--number", number, "--name", name])


def load(folder, number="12345"):
    return openpyxl.load_workbook(folder / pp.target_filename(number))


def today_cells(ws, col):
    return tuple(ws.cell(row=r, column=col).value for r in (3, 4, 5))


def test_creates_sheet_from_template(tmp_path):
    make_template(tmp_path)
    assert run(tmp_path) == 0
    wb = load(tmp_path)
    ws = wb.active
    assert ws["A3"].value == "Test Project"
    assert today_cells(ws, 4) == (TODAY.strftime("%d"), TODAY.strftime("%m"), TODAY.strftime("%y"))
    assert "Sheet 2" in wb.sheetnames


def test_appends_after_existing_dates(tmp_path):
    make_template(tmp_path, dates=[("01", "01", "20"), ("02", "02", "20")])
    run(tmp_path)
    ws = load(tmp_path).active
    assert today_cells(ws, 6)[0] == TODAY.strftime("%d")


def test_running_twice_same_day_reuses_column(tmp_path):
    make_template(tmp_path)
    run(tmp_path)
    run(tmp_path, name="Renamed")
    ws = load(tmp_path).active
    assert ws.cell(row=3, column=5).value is None
    assert ws["A3"].value == "Renamed"


def test_matches_integer_date_cells(tmp_path):
    make_template(tmp_path, dates=[(TODAY.day, TODAY.month, TODAY.year % 100)])
    run(tmp_path)
    assert load(tmp_path).active.cell(row=3, column=5).value is None


def test_prefers_placeholder_template(tmp_path):
    make_template(tmp_path, name="99999-SH-S-0001-Document Issue Sheet.xlsx", dates=[("01", "01", "20")])
    make_template(tmp_path)
    run(tmp_path)
    # Copied from the blank placeholder, so today's date lands in column D
    assert load(tmp_path).active.cell(row=3, column=4).value == TODAY.strftime("%d")


def test_missing_template_fails(tmp_path, capsys):
    assert run(tmp_path) == 1
    assert "No template" in capsys.readouterr().out


def test_strips_invalid_filename_chars(tmp_path):
    make_template(tmp_path)
    assert run(tmp_path, number="12/34?5") == 0
    assert (tmp_path / pp.target_filename("12345")).exists()


def test_prompts_when_no_args(tmp_path, monkeypatch):
    make_template(tmp_path)
    answers = iter(["", ""])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    assert pp.run(["--folder", str(tmp_path)]) == 0
    assert load(tmp_path, "00000").active["A3"].value == "Unnamed Project"


@pytest.mark.parametrize("filename, expected", [
    ("12345-DR-S-1001-P01-General Arrangement.pdf", ("12345-DR-S-1001", "P01", "General Arrangement")),
    ("12345-DR-S-1001-p02-Sections.pdf", ("12345-DR-S-1001", "P02", "Sections")),
    ("12345-DR-S-1001-C01.pdf", ("12345-DR-S-1001", "C01", "")),
    # Level code 'B1' in the number isn't mistaken for the revision
    ("12345-XX-B1-DR-S-1001-P03-Basement Plan.pdf", ("12345-XX-B1-DR-S-1001", "P03", "Basement Plan")),
    # Dashes in the title are kept
    ("12345-DR-S-1001-P01-Grid Line C1-C5.pdf", ("12345-DR-S-1001", "P01", "Grid Line C1-C5")),
])
def test_parse_drawing_name(tmp_path, filename, expected):
    drawing = pp.parse_drawing_name(tmp_path / filename)
    assert (drawing.number, drawing.revision, drawing.title) == expected


def test_parse_rejects_names_without_revision(tmp_path):
    assert pp.parse_drawing_name(tmp_path / "Specification.pdf") is None
    assert pp.parse_drawing_name(tmp_path / "12345-DR-S-1001.pdf") is None


def touch_pdfs(folder, *names):
    for name in names:
        (folder / name).write_bytes(b"%PDF-1.4")


def drawing_rows(ws):
    return [
        tuple(ws.cell(row=r, column=c).value for c in (1, 2, 4))
        for r in range(pp.DRAWING_START_ROW, ws.max_row + 1)
    ]


def test_lists_drawings_with_revisions(tmp_path):
    make_template(tmp_path)
    touch_pdfs(tmp_path,
               "12345-DR-S-1002-P01-Sections.pdf",
               "12345-DR-S-1001-P02-General Arrangement.pdf",
               "notes.txt")
    run(tmp_path)
    assert drawing_rows(load(tmp_path).active) == [
        ("12345-DR-S-1001", "General Arrangement", "P02"),
        ("12345-DR-S-1002", "Sections", "P01"),
    ]


def test_existing_drawings_keep_their_row(tmp_path):
    template = make_template(tmp_path, dates=[("01", "01", "20")])
    wb = openpyxl.load_workbook(template)
    ws = wb.active
    ws.cell(row=7, column=1, value="12345-DR-S-1005")
    ws.cell(row=7, column=2, value="My Own Title")
    ws.cell(row=7, column=4, value="P01")
    wb.save(template)

    touch_pdfs(tmp_path, "12345-DR-S-1005-P02-Pdf Title.pdf", "12345-DR-S-1001-P01-New.pdf")
    run(tmp_path)
    ws = load(tmp_path).active
    assert [ws.cell(row=r, column=c).value for r in (7, 8) for c in (1, 2, 4, 5)] == [
        "12345-DR-S-1005", "My Own Title", "P01", "P02",
        "12345-DR-S-1001", "New", None, "P01",
    ]


def test_newest_pdf_wins_for_duplicate_drawing(tmp_path):
    import os
    make_template(tmp_path)
    touch_pdfs(tmp_path, "12345-DR-S-1001-P02-GA.pdf", "12345-DR-S-1001-P01-GA.pdf")
    os.utime(tmp_path / "12345-DR-S-1001-P01-GA.pdf", (1, 1))
    run(tmp_path)
    assert drawing_rows(load(tmp_path).active) == [("12345-DR-S-1001", "GA", "P02")]


def test_unreadable_pdf_is_reported(tmp_path, capsys):
    make_template(tmp_path)
    touch_pdfs(tmp_path, "Specification.pdf")
    assert run(tmp_path) == 0
    assert "Skipped 'Specification.pdf'" in capsys.readouterr().out


def test_sheet_2_is_not_filled_with_drawings(tmp_path):
    make_template(tmp_path)
    touch_pdfs(tmp_path, "12345-DR-S-1001-P01-GA.pdf")
    run(tmp_path)
    assert load(tmp_path)["Sheet 2"].cell(row=7, column=1).value is None
