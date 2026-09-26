"""Record the PDF drawings in a folder as a new issue on the Document Issue Sheet.

Run it from the folder holding the drawing PDFs and the
"...SH-S-0001-Document Issue Sheet.xlsx" template. It creates
"<project number>-SH-S-0001-Document Issue Sheet.xlsx" from the template (if it
doesn't exist yet), writes the project name, adds a dated issue column, and
lists each PDF's drawing number, title and revision.

PDF names are read as <drawing number>-<revision>-<title>.pdf, for example
"12345-DR-S-1001-P01-General Arrangement.pdf". Running it again on the same
day reuses that day's column instead of adding a duplicate.
"""
import argparse
import datetime
import re
import shutil
import sys
from pathlib import Path
from typing import NamedTuple

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is not installed. Please run: pip install openpyxl")
    sys.exit(1)

TEMPLATE_SUFFIX = "SH-S-0001-Document Issue Sheet.xlsx"
FIRST_ISSUE_COL = 4  # Column D
DAY_ROW, MONTH_ROW, YEAR_ROW = 3, 4, 5
PROJECT_NAME_CELL = "A3"
DRAWING_START_ROW = 7
NUMBER_COL, TITLE_COL = 1, 2  # Columns A and B
SECOND_SHEET_NAME = "Sheet 2"

# Characters Windows doesn't allow in file names
INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*]')


# A revision code such as P01, C02 or T1
REVISION = re.compile(r"[A-Z]{1,2}\d{1,3}", re.IGNORECASE)


class Drawing(NamedTuple):
    number: str
    revision: str
    title: str
    path: Path


def parse_drawing_name(path):
    """Split '<number>-<revision>-<title>.pdf' into a Drawing, or return None.

    The revision is taken as the first revision-like part that follows a purely
    numeric part (the drawing's sequence number, e.g. '1001'), so codes such as
    a 'B1' level earlier in the number aren't mistaken for it.
    """
    parts = [part.strip() for part in path.stem.split("-")]
    candidates = [i for i in range(1, len(parts)) if REVISION.fullmatch(parts[i])]
    if not candidates:
        return None
    after_number = [i for i in candidates if parts[i - 1].isdigit()]
    rev_index = (after_number or candidates)[0]
    number = "-".join(parts[:rev_index])
    title = "-".join(parts[rev_index + 1:]).strip()
    return Drawing(number, parts[rev_index].upper(), title, path)


def find_drawings(folder):
    """Return (drawings, unreadable_names) for the PDFs in the folder.

    If a drawing appears more than once (e.g. P01 and P02 both present), the
    most recently modified file wins.
    """
    drawings, unreadable = {}, []
    for path in sorted(folder.iterdir()):
        if not path.is_file() or path.suffix.lower() != ".pdf":
            continue
        drawing = parse_drawing_name(path)
        if drawing is None:
            unreadable.append(path.name)
            continue
        key = drawing.number.upper()
        previous = drawings.get(key)
        if previous is None or path.stat().st_mtime > previous.path.stat().st_mtime:
            if previous is not None:
                print(f"Note: {drawing.number} has more than one PDF; using {path.name}")
            drawings[key] = drawing
        else:
            print(f"Note: {drawing.number} has more than one PDF; using {previous.path.name}")
    return sorted(drawings.values(), key=lambda d: d.number.upper()), unreadable


def write_drawings(sheet, col, drawings):
    """Write each drawing's revision into the issue column. Returns (added, updated).

    Drawings already listed in column A keep their row; new ones are added below
    the last listed drawing. A title is only filled in where the cell is empty,
    so hand-edited titles are kept.
    """
    rows = {}
    last_row = DRAWING_START_ROW - 1
    for row in range(DRAWING_START_ROW, sheet.max_row + 1):
        value = sheet.cell(row=row, column=NUMBER_COL).value
        if value is not None and str(value).strip():
            rows.setdefault(str(value).strip().upper(), row)
            last_row = row

    added = updated = 0
    for drawing in drawings:
        row = rows.get(drawing.number.upper())
        if row is None:
            last_row += 1
            row = last_row
            sheet.cell(row=row, column=NUMBER_COL, value=drawing.number)
            added += 1
        else:
            updated += 1
        title_cell = sheet.cell(row=row, column=TITLE_COL)
        if drawing.title and not title_cell.value:
            title_cell.value = drawing.title
        sheet.cell(row=row, column=col, value=drawing.revision)
    return added, updated


def target_filename(proj_number):
    return f"{proj_number}-{TEMPLATE_SUFFIX}"


def find_template(folder, exclude_name):
    """Return the template to copy, preferring placeholder names like 'xxxxx-...'.

    Files already numbered for a project (e.g. '12345-...') are only used as a
    fallback, and the project's own target file is never used as its template.
    """
    candidates = sorted(
        p for p in folder.glob(f"*{TEMPLATE_SUFFIX}")
        if p.name != exclude_name and not p.name.startswith("~$")  # skip Excel lock files
    )
    if not candidates:
        return None
    placeholders = [p for p in candidates if not p.name[0].isdigit()]
    return (placeholders or candidates)[0]


def _two_digits(value):
    """Normalise a header cell to a 2-digit string so '07', 7 and '7' compare equal."""
    if value is None:
        return None
    text = str(value).strip()
    return text.zfill(2) if text.isdigit() else text


def find_issue_column(sheet, date):
    """Return (column, already_present) for the given issue date.

    Scans right from column D. If a column already holds this date it's
    reused; otherwise the first column with an empty day cell is returned.
    """
    wanted = (date.strftime("%d"), date.strftime("%m"), date.strftime("%y"))
    col = FIRST_ISSUE_COL
    while sheet.cell(row=DAY_ROW, column=col).value is not None:
        existing = tuple(
            _two_digits(sheet.cell(row=row, column=col).value)
            for row in (DAY_ROW, MONTH_ROW, YEAR_ROW)
        )
        if existing == wanted:
            return col, True
        col += 1
    return col, False


def update_workbook(path, proj_name, date, drawings=()):
    """Write the project name, issue date and drawings into the workbook.

    Returns (column, already_present, added, updated).
    """
    wb = openpyxl.load_workbook(path)
    sheet = wb.active

    # Make sure the continuation sheet exists, copied before anything is written
    if SECOND_SHEET_NAME not in wb.sheetnames:
        wb.copy_worksheet(sheet).title = SECOND_SHEET_NAME

    col, already_present = find_issue_column(sheet, date)
    sheet.cell(row=DAY_ROW, column=col, value=date.strftime("%d"))
    sheet.cell(row=MONTH_ROW, column=col, value=date.strftime("%m"))
    sheet.cell(row=YEAR_ROW, column=col, value=date.strftime("%y"))
    sheet[PROJECT_NAME_CELL] = proj_name
    added, updated = write_drawings(sheet, col, drawings)

    wb.save(path)
    return col, already_present, added, updated


def prompt(message, default):
    answer = input(message).strip()
    return answer or default


def parse_args(argv):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--number", help="Revit project number, e.g. 12345")
    parser.add_argument("--name", help="Project name")
    parser.add_argument("--folder", type=Path, default=Path.cwd(),
                        help="Folder containing the PDFs and template (default: current folder)")
    return parser.parse_args(argv)


def run(argv=None):
    args = parse_args(argv)
    folder = args.folder.resolve()

    proj_number = args.number or prompt("Enter Revit Project Number (e.g., 12345): ", "00000")
    proj_number = INVALID_FILENAME_CHARS.sub("", proj_number).strip() or "00000"
    proj_name = args.name or prompt("Enter Project Name: ", "Unnamed Project")

    name = target_filename(proj_number)
    target_path = folder / name

    if not target_path.exists():
        template_path = find_template(folder, exclude_name=name)
        if template_path is None:
            print(f"Error: No template matching '*{TEMPLATE_SUFFIX}' found in {folder}")
            return 1
        shutil.copy(template_path, target_path)
        print(f"Created new issue sheet: {name} from template: {template_path.name}")

    drawings, unreadable = find_drawings(folder)
    for filename in unreadable:
        print(f"Skipped '{filename}': couldn't find a revision (expected e.g. 12345-DR-S-1001-P01-Title.pdf)")
    if not drawings:
        print(f"Warning: No drawing PDFs found in {folder}")

    try:
        col, already_present, added, updated = update_workbook(
            target_path, proj_name, datetime.date.today(), drawings)
    except PermissionError:
        print(f"Error: Could not save '{name}'. Please make sure it is CLOSED in Excel.")
        return 1
    except Exception as e:
        print(f"Error updating '{name}': {e}")
        return 1

    letter = openpyxl.utils.get_column_letter(col)
    if already_present:
        print(f"Today's issue column already existed (column {letter}); project name updated.")
    else:
        print(f"Added today's issue date in column {letter}.")
    print(f"Drawings: {added} added, {updated} already listed ({len(drawings)} issued).")
    print(f"Success! Saved: {target_path}")
    return 0


def main():
    code = run()
    # When run as a double-clicked .exe the console closes immediately, so wait first
    if getattr(sys, "frozen", False):
        input("\nPress Enter to close...")
    sys.exit(code)


if __name__ == "__main__":
    main()
